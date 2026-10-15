"""Consistency check of the built repository (app, data, exports) – run by CI before anything is merged.
Exit code 1 and a list of problems if something does not fit together, e.g. a stale build after a script change,
a missing file in the offline set, an unreplaced template placeholder or a key that slipped into the repo."""
import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)))
from base import BASE, FRAME, VERSION
import hashlib, json, os, re, subprocess, sys
import xml.etree.ElementTree as ET

ROOT = os.path.dirname(BASE)
problems = []
def bad(msg): problems.append(msg)
def rd(p): return open(os.path.join(ROOT, p), encoding='utf-8').read()
inside = lambda e, n: FRAME[0] - 200 <= e <= FRAME[2] + 200 and FRAME[1] - 200 <= n <= FRAME[3] + 200   # a few landmarks sit just outside the frame

# app shell: every placeholder of the template replaced
html = rd('index.html')
for ph in ('/*ISAPP*/', '<!--LEAFLETJS-->', '/*LEAFLET*/', 'const OFFLINE = null;', "const APP_VERSION = '';"):
    if ph in html: bad(f'index.html: Platzhalter {ph!r} nicht ersetzt')
if f"const APP_VERSION = '{VERSION}';" not in html: bad('index.html: App-Version passt nicht zu data/version.json')
m = re.search(r'const OFFLINE = (\{.*?\});', html)
off = json.loads(m.group(1)) if m else {'files': {}}
if not off['files']: bad('index.html: Offline-Liste leer')
sw = rd('sw.js')
if "const CORE = 'fc-core';" in sw: bad('sw.js: Cache-Name nicht ersetzt')
mk = re.search(r'const KEEP = (\[.*?\]);', sw); keep = set(json.loads(mk.group(1))) if mk else set()
if keep != {f'{p}?v={h}' for p, h in off['files'].items()}: bad('sw.js und index.html nennen unterschiedliche Offline-Dateien')
# offline files exist and their hashes are current (a stale hash would serve an old file forever)
for p, h in off['files'].items():
    fp = os.path.join(ROOT, p)
    if not os.path.exists(fp): bad(f'Offline-Datei fehlt: {p}'); continue
    if hashlib.md5(open(fp, 'rb').read()).hexdigest()[:10] != h: bad(f'Offline-Datei veraltet (Hash): {p} – build_site.py neu laufen lassen')

# data.json
d = json.loads(rd('data.json'))
if d.get('version') != VERSION: bad(f"data.json: Version {d.get('version')} statt {VERSION}")
ids = [i['id'] for i in d['items']] + [p['id'] for p in d['paths']]
dup = {i for i in ids if ids.count(i) > 1}
if dup: bad(f'data.json: doppelte IDs {sorted(dup)[:10]}')
for i in d['items']:
    if not (isinstance(i.get('e'), (int, float)) and isinstance(i.get('nn'), (int, float)) and inside(i['e'], i['nn'])): bad(f"data.json: {i['id']} ohne gültige Lage")
for p in d['paths']:
    if len(p.get('c', [])) < 2 or not all(inside(*c) for c in p['c']): bad(f"data.json: Pfad {p['id']} ungültig")
if any(not k.startswith('d-') for k in d.get('feld', {})): bad('data.json: feld enthält Namen statt Geräte-IDs')
for a in d.get('sheets', []):
    if len(a.get('bbox', [])) != 4: bad(f"data.json: Gebiet {a.get('id')} ohne bbox")

# inbox config
try:
    bk = json.loads(rd('briefkasten.json'))
    if not re.match(r'^[\w.-]+/[\w.-]+$', bk.get('repo', '')) or not bk.get('zweig') or not bk.get('ordner'): bad('briefkasten.json unvollständig')
except Exception as e: bad(f'briefkasten.json unlesbar ({e})')

# PDF data for the app
lay = json.loads(rd('karten/layers.json'))
for s in lay['sheets']:
    for f in s['files'].values():
        if not os.path.exists(os.path.join(ROOT, f)): bad(f'karten/layers.json nennt fehlende Datei {f}')

# OL export
try:
    x = ET.parse(os.path.join(BASE, 'out/ol/Wangerooge_West_Posten.xml'))
    ns = {'i': 'http://www.orienteering.org/datastandard/3.0'}
    n_ctl = len(x.findall('.//i:Control[@type="Control"]', ns)); n_sure = sum(1 for i in d['items'] if i['s'] == 'sicher')
    if n_ctl != n_sure: bad(f'IOF-XML: {n_ctl} Posten, aber {n_sure} sichere Objekte in data.json – Export veraltet')
except FileNotFoundError: bad('out/ol/Wangerooge_West_Posten.xml fehlt')
except ET.ParseError as e: bad(f'IOF-XML kaputt: {e}')

# Excel list matches the data
try:
    import openpyxl
    ws = openpyxl.load_workbook(os.path.join(BASE, 'out/Wangerooge_West_Landmarken_v1.xlsx'), read_only=True)['Landmarken']
    if ws.max_row - 1 != len(d['items']): bad(f'Excel: {ws.max_row - 1} Zeilen, data.json: {len(d["items"])} Objekte')
except ImportError: print('Hinweis: openpyxl fehlt, Excel nicht geprüft')

# no keys in the repository
TOKEN = re.compile(rb'github_pat_[A-Za-z0-9_]{30,}|gh[pousr]_[A-Za-z0-9]{30,}')
files = subprocess.run(['git', 'ls-files', '-z'], cwd=ROOT, capture_output=True).stdout.split(b'\0')
for f in files:
    f = f.decode()
    if not f or f.endswith(('.png', '.jpg', '.pdf', '.tif', '.kmz', '.xlsx', '.pkl', '.npy', '.woff2', '.ttf', '.otf')): continue
    fp = os.path.join(ROOT, f)
    if os.path.isfile(fp) and os.path.getsize(fp) < 50_000_000 and TOKEN.search(open(fp, 'rb').read()): bad(f'Schlüssel im Repo: {f}')

if problems:
    print('Prüfung fehlgeschlagen:'); [print(' -', p) for p in problems]; sys.exit(1)
print(f'Prüfung ok: {len(d["items"])} Objekte, {len(d["paths"])} Pfade, {len(off["files"])} Offline-Dateien, Version {VERSION}')
