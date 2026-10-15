import os as _os, sys as _sys; _sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__))); from base import BASE, VERSION, VERSION_DATE, de_date, FRAME, OSM_DATE, DGM_YEAR, DOP_DATE
import csv, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
rows = list(csv.DictReader(open(BASE + '/out/landmarken_v1.csv', encoding='utf-8')))
paths = json.load(open(BASE + '/web/site/data.json'))['paths']
A = lambda **k: Font(name='Arial', **k)
H = PatternFill('solid', fgColor='2F2A36'); G = PatternFill('solid', fgColor='EEEAE2')
wb = Workbook(); ws = wb.active; ws.title = 'Landmarken'
ws.append(['ID', 'Kategorie', 'Typ', 'Name', 'Status', 'Quellen', 'Breite (WGS84)', 'Länge (WGS84)', 'UTM32 Ost (m)', 'UTM32 Nord (m)', 'Zone', 'Notiz', 'Vor Ort (Feldcheck)', 'OSM-ID'])
for r in sorted(rows, key=lambda r: (r['kategorie'], r['id'])):
    ws.append([r['id'], r['kategorie'], r['typ'], r['name'], r['status'], r['quellen'], float(r['lat']), float(r['lon']), float(r['utm_e']), float(r['utm_n']), r['zone'], r['notiz'], r.get('vor_ort', ''), r['osm']])
for c in ws[1]: c.font = A(bold=True, color='FFFFFF'); c.fill = H; c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = A(size=10)
    row[6].number_format = '0.000000'; row[7].number_format = '0.000000'; row[8].number_format = '0.0'; row[9].number_format = '0.0'
for k, v in {'A': 9, 'B': 15, 'C': 18, 'D': 30, 'E': 9, 'F': 34, 'G': 12, 'H': 12, 'I': 12, 'J': 12, 'K': 14, 'L': 64, 'M': 40, 'N': 14}.items(): ws.column_dimensions[k].width = v
ws.freeze_panes = 'B2'; ws.auto_filter.ref = f"A1:N{ws.max_row}"
ws.conditional_formatting.add(f"A2:N{ws.max_row}", FormulaRule(formula=['$E2="prüfen"'], font=Font(name='Arial', color='8A8A8A', size=10)))
wp = wb.create_sheet('Pfade prüfen'); wp.append(['ID', 'Wegtyp', 'Info', 'Quelle'])
for p in paths: wp.append([p['id'], p['t'], ' '.join(x for x in (p.get('f'), p.get('w')) if x), p['q']])
for c in wp[1]: c.font = A(bold=True, color='FFFFFF'); c.fill = H
for row in wp.iter_rows(min_row=2):
    for c in row: c.font = A(size=10)
for k, v in {'A': 14, 'B': 12, 'C': 60, 'D': 32}.items(): wp.column_dimensions[k].width = v
wp.freeze_panes = 'A2'
wo = wb.create_sheet('Übersicht', 0)
wo['A1'] = 'Wangerooge West – Landmarkenliste'; wo['A1'].font = A(bold=True, size=16)
wo['A2'] = f'Version {VERSION} vom {de_date(VERSION_DATE)} · Kartenrahmen UTM32 E {FRAME[0]}–{FRAME[2]}, N {FRAME[1]}–{FRAME[3]} · Start: Inselheim Rüstringen mit Sternwarte'; wo['A2'].font = A(size=10, color='555555')
wo['A4'] = 'Status „sicher“ = mindestens zwei unabhängige Quellen. „prüfen“ = nur eine Quelle; in der Karte normal gezeichnet, in der Feldcheck-App als weißer Marker. IDs sind ab Version 1.1 stabil.'; wo['A4'].font = A(size=10)
wo['A6'], wo['B6'], wo['C6'], wo['D6'] = 'Kategorie', 'sicher', 'prüfen', 'gesamt'
for c in wo[6]: c.font = A(bold=True, color='FFFFFF'); c.fill = H
cats = []
for r in rows:
    if r['kategorie'] not in cats: cats.append(r['kategorie'])
n = len(rows) + 1
for i, k in enumerate(cats):
    rr = 7 + i
    wo.cell(rr, 1, k)
    wo.cell(rr, 2, f'=COUNTIFS(Landmarken!$B$2:$B${n},$A{rr},Landmarken!$E$2:$E${n},"sicher")')
    wo.cell(rr, 3, f'=COUNTIFS(Landmarken!$B$2:$B${n},$A{rr},Landmarken!$E$2:$E${n},"prüfen")')
    wo.cell(rr, 4, f'=B{rr}+C{rr}')
    for c in range(1, 5): wo.cell(rr, c).font = A(size=10)
rp = 7 + len(cats)
wo.cell(rp, 1, 'Pfade (Luftbild unklar)'); wo.cell(rp, 2, 0); wo.cell(rp, 3, f"=COUNTA('Pfade prüfen'!A2:A{len(paths) + 1})"); wo.cell(rp, 4, f'=B{rp}+C{rp}')
for c in range(1, 5): wo.cell(rp, c).font = A(size=10)
wo.cell(rp, 2).font = A(size=10, color='0000FF')
wo.cell(rp, 2).comment = Comment('Alle übrigen Pfade sind sicher (befestigt oder im Luftbild sichtbar); hier zählen nur die zu prüfenden.', 'Claude')
rt = rp + 1
wo.cell(rt, 1, 'Summe')
for c in (2, 3, 4): wo.cell(rt, c, f'=SUM({get_column_letter(c)}7:{get_column_letter(c)}{rp})')
for c in range(1, 5): wo.cell(rt, c).font = A(bold=True); wo.cell(rt, c).fill = G
s = rt + 2
src = [('Quellen', ''), (f'OpenStreetMap (Abruf {de_date(OSM_DATE)})', '© OpenStreetMap-Mitwirkende, ODbL – Wege, Gebäude, Bänke, POIs, Ruhezone I/32'),
       (f'LGLN DGM1 (Befliegung {DGM_YEAR})', 'LGLN (2026), CC BY 4.0 – Höhenlinien, Gruben/Trichter, Aussichtsdüne'),
       (f'LGLN DOP20 (Befliegung {de_date(DOP_DATE)})', 'LGLN (2026), CC BY 4.0 – Bank-/Trichterabgleich, Sandflächen, Pfadsichtbarkeit, Lage Sternwarte'),
       ('aloeus.de, Wikipedia, TracesOfWar, gruppenhaus.de', 'Zweitquellen für Sternwarte, Westturm, Leuchtturm, Kriegsgräberstätte, Heime'),
       ('Hinweis', 'Ruhezonen-Grenze stammt aus OSM (aus der Seekarte übernommen); vor Ort an den Schildern prüfen.')]
for i, (a, b) in enumerate(src): wo.cell(s + i, 1, a).font = A(bold=(i == 0), size=10); wo.cell(s + i, 2, b).font = A(size=10)
wo.column_dimensions['A'].width = 34
for c in 'BCD': wo.column_dimensions[c].width = 12
# same data -> byte-identical file: fixed document times and zip entry times (openpyxl stamps 'now' on save)
import datetime as _dt, io as _io, re as _re, zipfile as _zf
wb.properties.created = _dt.datetime.fromisoformat(VERSION_DATE + 'T12:00:00')
buf = _io.BytesIO(); wb.save(buf); src = _zf.ZipFile(buf)
ZT = tuple(int(v) for v in VERSION_DATE.split('-')) + (12, 0, 0)
with _zf.ZipFile(BASE + '/out/Wangerooge_West_Landmarken_v1.xlsx', 'w', _zf.ZIP_DEFLATED) as z:
    for zi in src.infolist():
        data = src.read(zi.filename)
        if zi.filename == 'docProps/core.xml':
            data = _re.sub(rb'(<dcterms:modified[^>]*>)[^<]*', rb'\g<1>' + VERSION_DATE.encode() + b'T12:00:00Z', data)
        n = _zf.ZipInfo(zi.filename, ZT); n.compress_type = _zf.ZIP_DEFLATED; z.writestr(n, data)
print('ok', len(rows), len(paths))
